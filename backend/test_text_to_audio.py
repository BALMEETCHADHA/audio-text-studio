#!/usr/bin/env python3
"""
Tests for the Text -> Audio side and the accuracy changes — run from backend/ with:

    python3 test_text_to_audio.py

No internet or API key is needed: the voice engines and the AI model are replaced by fakes, so these
check OUR logic (voice list, text splitting, caching, zip bundles, file reading, the spelling-fix
safety checks, summary length). Whether Microsoft's voices sound right is something only you can
check by pressing the ▶ buttons.
"""
import sys
import tempfile
import traceback
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
failures = 0


def check(name, fn):
    global failures
    try:
        fn()
        print(f"  PASS  {name}")
    except Exception:
        failures += 1
        print(f"  FAIL  {name}")
        print("        " + traceback.format_exc().strip().replace("\n", "\n        "))


from config import settings  # noqa: E402
from pathlib import Path  # noqa: E402
from services import tts, tts_voices  # noqa: E402

settings.storage_dir = tempfile.mkdtemp()
settings.storage_path  # creates the folders

print("\n1. The voice list")


def voice_list_matches_the_request():
    v = tts_voices.VOICES
    assert len(v) >= 20, len(v)
    india = [x for x in v if x.group == "India"]
    assert len(india) == 8
    assert sum(x.gender == "Male" for x in india) == 4 and sum(x.gender == "Female" for x in india) == 4
    assert [x.id for x in v[:8]] == [f"india-{i}" for i in range(1, 9)], "India must be at the top"
    for group in ("United States", "Canada", "Australia"):
        assert any(x.group == group for x in v), group
    for group in ("Japan", "China"):
        assert sum(x.group == group for x in v) >= 2, group
    assert len({x.id for x in v}) == len(v), "ids must be unique"
    assert tts_voices.DEFAULT_VOICE_ID == "india-1"


def labels_are_the_requested_style():
    labels = [x.label for x in tts_voices.VOICES]
    assert "India 1" in labels and "India 2" in labels and "US English 1" in labels and "Canadian English 1" in labels


check("20+ voices; India 1-8 first (4 female, 4 male); US/Canada/Australia; 2 each of Japan and China", voice_list_matches_the_request)
check("labels look like 'India 1', 'US English 1', 'Canadian English 1'", labels_are_the_requested_style)

print("\n2. Preparing text for speech")


def splitting_keeps_everything_and_respects_the_limit():
    text = " ".join(f"This is sentence number {i}." for i in range(400))
    chunks = tts.split_text(text, limit=500)
    assert all(len(c) <= 500 for c in chunks) and len(chunks) > 3
    assert " ".join(chunks).split() == text.split(), "no word may be lost or reordered"
    giant = "word " * 2000  # one endless sentence
    assert all(len(c) <= 300 for c in tts.split_text(giant, limit=300))


def markdown_marks_are_not_read_aloud():
    assert "*" not in tts.clean_for_speech("**Bold** and `code` and # heading")


check("long text is cut at sentence ends without losing words", splitting_keeps_everything_and_respects_the_limit)
check("markdown symbols are removed before speaking", markdown_marks_are_not_read_aloud)

print("\n3. Making audio (voice engine faked)")
calls = []


async def fake_edge(chunks, voice_name):
    calls.append((voice_name, len(chunks)))
    return [f"MP3[{voice_name}]".encode() for _ in chunks]


tts._edge_all = fake_edge


def audio_is_cached_per_voice_and_text():
    calls.clear()
    a = tts.synthesize_to_file("Hello there world.", "india-1")
    b = tts.synthesize_to_file("Hello there world.", "india-1")
    c = tts.synthesize_to_file("Hello there world.", "us-1")
    assert a == b and a != c and a.suffix == ".mp3" and a.read_bytes().startswith(b"MP3[en-IN-Neerja")
    assert len(calls) == 2, "the second request for the same voice+text must come from disk"


def previews_are_made_once():
    calls.clear()
    p = tts.preview_path("japan-1")
    assert p.exists() and tts.preview_path("japan-1") == p and len(calls) == 1
    try:
        tts.preview_path("nope")
        raise AssertionError("unknown voice")
    except KeyError:
        pass


def bundle_single_is_mp3_several_is_zip():
    path, fmt = tts.build_audio_bundle("Some text to read.", ["india-1"], "whole-text")
    assert fmt == "mp3" and path.suffix == ".mp3"
    seen = []
    path, fmt = tts.build_audio_bundle("Some text to read.", ["india-1", "us-1", "china-2"], "summary", lambda d, t: seen.append((d, t)))
    assert fmt == "zip" and seen == [(1, 3), (2, 3), (3, 3)]
    names = zipfile.ZipFile(path).namelist()
    assert len(names) == 3 and all(n.endswith(".mp3") and n.startswith("summary-") for n in names), names


def failing_engine_gives_a_readable_error():
    async def boom(chunks, name):
        raise RuntimeError("blocked")

    saved, tts._edge_all = tts._edge_all, boom
    settings.tts_engine = "edge"
    try:
        tts.synthesize("Text that fails.", "india-1")
        raise AssertionError("should have raised")
    except tts.TTSError as exc:
        assert "Couldn't create the audio" in str(exc)
    finally:
        tts._edge_all = saved
        settings.tts_engine = "auto"


def empty_text_is_refused():
    try:
        tts.synthesize("  ** ", "india-1")
        raise AssertionError("should have raised")
    except tts.TTSError:
        pass


check("same text + voice is made once and reused; other voices get their own file", audio_is_cached_per_voice_and_text)
check("voice samples are generated once", previews_are_made_once)
check("one accent -> MP3, several -> ZIP of MP3s", bundle_single_is_mp3_several_is_zip)
check("an engine failure becomes a readable message", failing_engine_gives_a_readable_error)
check("nothing to read is refused", empty_text_is_refused)

print("\n4. Reading uploaded files")


def reads_txt_and_docx():
    from docx import Document
    from services import text_extract

    d = Path(tempfile.mkdtemp())
    (d / "a.txt").write_text("Plain   text\n\n\n\nfile here.", encoding="utf-8")
    assert text_extract.extract_text(d / "a.txt", "a.txt") == "Plain text\n\nfile here."
    doc = Document()
    doc.add_paragraph("First paragraph.")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Cell A", "Cell B"
    doc.save(d / "b.docx")
    out = text_extract.extract_text(d / "b.docx", "b.docx")
    assert "First paragraph." in out and "Cell A | Cell B" in out


def unsupported_and_empty_files_get_clear_errors():
    from services import text_extract

    d = Path(tempfile.mkdtemp())
    (d / "x.exe").write_bytes(b"MZ")
    (d / "e.txt").write_text("   ")
    (d / "o.doc").write_bytes(b"x")
    for name, word in (("x.exe", "PDF"), ("e.txt", "No readable text"), ("o.doc", "docx")):
        try:
            text_extract.extract_text(d / name, name)
            raise AssertionError(name)
        except text_extract.ExtractionError as exc:
            assert word in str(exc), (name, str(exc))


def pdf_with_text_is_read():
    try:
        from fpdf import FPDF
        import pypdf  # noqa: F401
    except ImportError:
        print("        (skipped: fpdf2/pypdf not installed here)")
        return
    from services import text_extract

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, "The quick brown fox jumps over the lazy dog, again and again.")
    p = Path(tempfile.mkdtemp()) / "t.pdf"
    pdf.output(str(p))
    assert "quick brown fox" in text_extract.extract_text(p, "t.pdf")


check("txt and docx (paragraphs + tables) are read", reads_txt_and_docx)
check("unsupported / empty / old .doc files give readable errors", unsupported_and_empty_files_get_clear_errors)
check("a PDF with a text layer is read", pdf_with_text_is_read)

print("\n5. Smart spelling fix — must never damage the transcript")
from services import analysis, correction  # noqa: E402

ORIGINAL = "I interned at Edinet Foundation under ECTE and built a BILSTM model for stock forecasting with YFinance and GARCH."


def accepts_small_repairs_only():
    fixed = ORIGINAL.replace("Edinet", "Edunet").replace("ECTE", "AICTE")
    assert correction.accept_correction(ORIGINAL, fixed)
    assert not correction.accept_correction(ORIGINAL, "Totally different text now.")
    assert not correction.accept_correction(ORIGINAL, ORIGINAL + " And a whole extra invented sentence appended.")
    assert not correction.accept_correction(ORIGINAL, "")


def lists_what_changed():
    fixed = ORIGINAL.replace("Edinet", "Edunet").replace("ECTE", "AICTE")
    changes = correction.list_changes(ORIGINAL, fixed)
    assert {"from": "Edinet", "to": "Edunet"} in changes and {"from": "ECTE", "to": "AICTE"} in changes


def applies_good_fix_and_rejects_a_rewrite():
    good = ORIGINAL.replace("Edinet", "Edunet").replace("ECTE", "AICTE")
    analysis.run_json_prompt = lambda prompt: {"corrected": good}
    text, changes = correction.correct_transcript(ORIGINAL, "Edunet Foundation, AICTE", None)
    assert text == good and len(changes) == 2
    analysis.run_json_prompt = lambda prompt: {"corrected": "I did some internships. Many things were learned."}
    text, changes = correction.correct_transcript(ORIGINAL)
    assert text == ORIGINAL and changes == [], "a rewrite must be thrown away"


def model_failure_keeps_the_transcript():
    def boom(prompt):
        raise RuntimeError("rate limited")

    analysis.run_json_prompt = boom
    text, changes = correction.correct_transcript(ORIGINAL)
    assert text == ORIGINAL and changes == []


check("only small word repairs are accepted", accepts_small_repairs_only)
check("the changed words are listed (Edinet -> Edunet)", lists_what_changed)
check("a good fix is applied; a rewrite is rejected", applies_good_fix_and_rejects_a_rewrite)
check("an AI failure never costs the transcript", model_failure_keeps_the_transcript)

print("\n6. Summary length (30 / 50 / 70 / 100 / custom)")
import worker  # noqa: E402


def summary_length_is_clamped_not_snapped():
    assert [worker.clamp_summary_words(x) for x in (30, 50, 70, 100)] == [30, 50, 70, 100]
    assert worker.clamp_summary_words(150) == 150 and worker.clamp_summary_words("85") == 85
    assert worker.clamp_summary_words(3) == 10 and worker.clamp_summary_words(99999) == 500
    assert worker.clamp_summary_words("abc") == 70


def local_summary_respects_a_short_target():
    from services.local_analysis import analyze_transcript_local

    text = " ".join(f"This is sentence number {i} and it has quite a few extra words inside it for length." for i in range(20))
    r = analyze_transcript_local(text, summary_words=30, include_key_points=False, include_highlights=False)
    assert len(r["summary"].split()) <= 36, len(r["summary"].split())


check("any number is accepted, kept between 10 and 500", summary_length_is_clamped_not_snapped)
check("the offline summary stays near a 30-word target", local_summary_respects_a_short_target)

print("\n7. Downloads in the new order")


def docx_and_image_follow_page_order():
    from docx import Document
    from services import export_docx, export_image

    out = export_docx.build_full_report_docx(
        transcript="The whole text.", summary="Short summary.", key_points=["Point"],
        dates_mentioned=["May 2026"], links_mentioned=[], notable_terms=["Whisper"],
        ai_insights="Some insight.", title="Text and Audio Report", transcript_label="Full Text",
    )
    heads = [p.text for p in Document(str(out)).paragraphs if p.style.name.startswith(("Heading", "Title"))]
    order = ["Full Text", "Summary", "Key Points", "AI Insights", "Highlights"]
    idx = [heads.index(h) for h in order]
    assert idx == sorted(idx), heads

    img = export_image.build_summary_image(
        title="T", summary="S " * 40, key_points=["a", "b"], transcript="word " * 3000,
        ai_insights="insight " * 30, dates=["x"], links=["y.com"], notable_terms=["z"],
    )
    from PIL import Image

    with Image.open(img) as im:
        assert im.height > 800 and im.height <= export_image.MAX_HEIGHT


check("Word report is ordered: text, summary, key points, AI insights, highlights; image handles huge text", docx_and_image_follow_page_order)

print("\n8. Summary + whole text in one MP3")


def combined_audio_joins_summary_then_text():
    calls.clear()
    path, fmt = tts.build_audio_bundle(["Summary. Short.", "Full text. Long long text."], ["india-1"], "summary-and-text")
    assert fmt == "mp3"
    data = path.read_bytes()
    assert data.count(b"MP3[en-IN-Neerja") == 2, "two spoken parts, one after the other"
    again, _ = tts.build_audio_bundle(["Summary. Short.", "Full text. Long long text."], ["india-1"], "summary-and-text")
    assert again == path
    p2, f2 = tts.build_audio_bundle(["Summary. Short.", "Full text. Long long text."], ["india-1", "uk-1"], "summary-and-text")
    assert f2 == "zip" and len(zipfile.ZipFile(p2).namelist()) == 2


check("two texts become one MP3 per accent (ZIP for several accents)", combined_audio_joins_summary_then_text)

print("\n9. Speaker labels")
from services import diarization  # noqa: E402


def speaker_turns_are_merged_and_numbered_by_first_appearance():
    segs = [
        {"text": "Hello everyone.", "speaker": "Speaker 07"},
        {"text": "Thanks for joining.", "speaker": "Speaker 07"},
        {"text": "Happy to be here.", "speaker": "Speaker 02"},
        {"text": "Let's begin.", "speaker": None},   # no label: carried over from the previous turn
        {"text": "Great.", "speaker": "Speaker 07"},
    ]
    out = diarization.speaker_transcript(segs)
    assert out == (
        "Speaker 1: Hello everyone. Thanks for joining.\n\n"
        "Speaker 2: Happy to be here. Let's begin.\n\n"
        "Speaker 1: Great."
    ), out
    assert diarization.speaker_transcript([{"text": "x"}]) is None


def worker_skips_labels_politely_when_not_set_up():
    warnings = []
    segs, labeled = worker._speaker_labels(Path("x.wav"), [{"text": "hi"}], True, [], warnings)
    assert labeled is None and warnings and "not set up" in warnings[0]
    segs, labeled = worker._speaker_labels(Path("x.wav"), [{"text": "hi"}], False, [], warnings)
    assert labeled is None and len(warnings) == 1


def spelling_fixes_carry_into_the_labelled_text():
    saved = diarization.diarization_available, diarization.label_speakers
    diarization.diarization_available = lambda: True
    diarization.label_speakers = lambda path, segs: [{**s, "speaker": "A"} for s in segs]
    try:
        _, labeled = worker._speaker_labels(
            Path("x.wav"), [{"text": "I joined Edinet Foundation."}], True, [{"from": "Edinet", "to": "Edunet"}], []
        )
    finally:
        diarization.diarization_available, diarization.label_speakers = saved
    assert labeled == "Speaker 1: I joined Edunet Foundation.", labeled


check("turns are merged and numbered Speaker 1, 2, 3 by who speaks first", speaker_turns_are_merged_and_numbered_by_first_appearance)
check("asking for labels on a server without them gives a note, not a crash", worker_skips_labels_politely_when_not_set_up)
check("spelling fixes also apply inside the speaker-labelled text", spelling_fixes_carry_into_the_labelled_text)

print("\n10. Access keys and the daily limit")
import auth  # noqa: E402
import jobstore  # noqa: E402


class FakeRequest:
    def __init__(self, ip="1.1.1.1", forwarded=None):
        self.headers = {"x-forwarded-for": forwarded} if forwarded else {}
        self.client = type("C", (), {"host": ip})()
        self.state = type("S", (), {})()


def keys_can_be_named_or_plain():
    settings.access_key = "ravi=k7Hq2xPz, priya=Vb83nLmQ ,plainkey"
    names = [n for n, _ in auth.parse_keys()]
    assert names[:2] == ["ravi", "priya"] and names[2].startswith("key-") and auth.auth_enabled()
    settings.access_key = ""
    assert not auth.auth_enabled()


def each_key_signs_in_as_its_owner_and_wrong_keys_are_refused():
    settings.access_key = "ravi=k7Hq2xPz,priya=Vb83nLmQ"
    try:
        r = FakeRequest()
        auth.require_api_key(r, "Vb83nLmQ")
        assert r.state.user == "priya" and auth.owner_of(r) == "user:priya"
        for bad in (None, "", "nope"):
            try:
                auth.require_api_key(FakeRequest(), bad)
                raise AssertionError("should be refused")
            except Exception as exc:
                assert getattr(exc, "status_code", None) == 401, exc
    finally:
        settings.access_key = ""


def ten_wrong_keys_lock_the_visitor_out_for_a_while():
    settings.access_key = "ravi=k7Hq2xPz"
    auth._failures.clear()
    try:
        codes = []
        for _ in range(12):
            try:
                auth.require_api_key(FakeRequest(ip="9.9.9.9"), "wrong")
            except Exception as exc:
                codes.append(exc.status_code)
        assert codes[:10] == [401] * 10 and codes[10:] == [429, 429], codes
        auth.require_api_key(FakeRequest(ip="8.8.8.8"), "k7Hq2xPz")  # someone else is unaffected
        # even the right key is paused during the lockout
        try:
            auth.require_api_key(FakeRequest(ip="9.9.9.9"), "k7Hq2xPz")
        except Exception as exc:
            assert exc.status_code == 429
    finally:
        settings.access_key = ""
        auth._failures.clear()


def real_address_is_read_behind_a_proxy():
    r = FakeRequest(ip="10.0.0.1", forwarded="203.0.113.7, 10.0.0.1")
    settings.trust_proxy = True
    assert auth.client_ip(r) == "203.0.113.7"
    settings.trust_proxy = False
    assert auth.client_ip(r) == "10.0.0.1"
    settings.trust_proxy = True


def daily_limit_counts_jobs_per_owner():
    jobstore.init_db()
    a = f"user:test-{id(object())}"
    b = a + "-other"
    for _ in range(3):
        jobstore.create_job("full", "x.mp3", owner=a)
    jobstore.create_job("full", "y.mp3", owner=b)
    assert jobstore.count_jobs_last_day(a) == 3 and jobstore.count_jobs_last_day(b) == 1


check("ACCESS_KEY accepts named keys, plain keys, or nothing", keys_can_be_named_or_plain)
check("each key signs in as its owner; missing/wrong keys get 401", each_key_signs_in_as_its_owner_and_wrong_keys_are_refused)
check("10 wrong keys in a row pause that visitor (429); others are unaffected", ten_wrong_keys_lock_the_visitor_out_for_a_while)
check("the visitor's real IP is read from X-Forwarded-For behind a proxy", real_address_is_read_behind_a_proxy)
check("jobs are counted per owner over the last 24 hours", daily_limit_counts_jobs_per_owner)

print(
    "\nAll text-to-audio checks passed.\n"
    if failures == 0
    else f"\n{failures} check(s) FAILED — fix before shipping.\n"
)
sys.exit(1 if failures else 0)
