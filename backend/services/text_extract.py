"""
Pulls plain text out of a file the person uploads on the Text -> Audio page:
PDF, Word (.docx), an image (read with OCR) or a .txt file.

* PDF: the text layer is read directly. A scanned PDF has none, so its pages are drawn to pictures
  and read with OCR instead (needs the Tesseract program on the server — see README).
* Image: OCR with Tesseract (English).
* Every failure is raised as ExtractionError with a sentence the person can act on.
"""
import re
from pathlib import Path


class ExtractionError(Exception):
    """A readable reason the text couldn't be taken from the file."""


_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
_MAX_OCR_PAGES = 30
_TESSERACT_HELP = (
    "Reading text from pictures needs the free Tesseract OCR program on the server "
    "(Windows: https://github.com/UB-Mannheim/tesseract/wiki — Linux: apt install tesseract-ocr — "
    "Mac: brew install tesseract) plus `pip install pytesseract`."
)


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _ocr_image(image) -> str:
    try:
        import pytesseract
    except ImportError as exc:
        raise ExtractionError(_TESSERACT_HELP) from exc
    try:
        return pytesseract.image_to_string(image, lang="eng")
    except pytesseract.TesseractNotFoundError as exc:  # type: ignore[attr-defined]
        raise ExtractionError(_TESSERACT_HELP) from exc
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Couldn't read the text in that picture: {exc}") from exc


def _from_image(path: Path) -> str:
    try:
        from PIL import Image, ImageOps
    except ImportError as exc:  # pragma: no cover — Pillow is a hard requirement of the app
        raise ExtractionError("Image support (Pillow) is not installed.") from exc
    try:
        with Image.open(path) as img:
            img = ImageOps.exif_transpose(img).convert("L")
            if max(img.size) < 1200:  # small screenshots read far better when enlarged
                img = img.resize((img.width * 2, img.height * 2))
            text = _ocr_image(img)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("That image couldn't be opened. Try a PNG or JPG.") from exc
    return text


def _from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF support isn't installed on the server (pip install pypdf).") from exc

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                pass
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("That PDF couldn't be read. It may be damaged or password-protected.") from exc

    text = _tidy("\n\n".join(pages))
    if len(re.sub(r"\s", "", text)) >= 20:
        return text

    # No text layer: probably a scan. Draw each page and read it with OCR.
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise ExtractionError(
            "This PDF looks like a scan (it has no selectable text). Reading scans needs "
            "`pip install pypdfium2 pytesseract` and the Tesseract program on the server."
        ) from exc
    try:
        pdf = pdfium.PdfDocument(str(path))
        out = []
        for index in range(min(len(pdf), _MAX_OCR_PAGES)):
            image = pdf[index].render(scale=2).to_pil().convert("L")
            out.append(_ocr_image(image))
        return "\n\n".join(out)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Couldn't read the scanned pages of that PDF: {exc}") from exc


def _from_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("Word support (python-docx) is not installed.") from exc
    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("That Word file couldn't be opened. Save it as .docx and try again.") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(path: Path, filename: str | None = None) -> str:
    """Return the text found in `path`. `filename` (the original name) decides the type."""
    ext = Path(filename or path.name).suffix.lower()
    if ext == ".pdf":
        text = _from_pdf(path)
    elif ext == ".docx":
        text = _from_docx(path)
    elif ext == ".doc":
        raise ExtractionError("Old .doc files aren't supported — open it in Word and 'Save as' .docx, or upload a PDF.")
    elif ext in _IMAGE_EXT:
        text = _from_image(path)
    elif ext in {".txt", ".md", ".text", ""}:
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            raise ExtractionError("That text file couldn't be read.") from exc
    else:
        raise ExtractionError("Please upload a PDF, a Word (.docx) file or an image (PNG / JPG).")

    text = _tidy(text)
    if not text:
        raise ExtractionError("No readable text was found in that file.")
    return text
