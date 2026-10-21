"""
Word document versions of the results — same content as the PDF, in a format
that's easier for most people to open and edit further.
"""
import uuid
from pathlib import Path

from docx import Document

from config import settings


def _add_bullets(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _save(doc: Document) -> Path:
    out_path = settings.storage_path / "reports" / f"{uuid.uuid4().hex}.docx"
    doc.save(str(out_path))
    return out_path


def build_full_report_docx(
    transcript: str,
    summary: str | None,
    key_points: list[str],
    dates_mentioned: list[str],
    links_mentioned: list[str],
    notable_terms: list[str],
    ai_insights: str | None,
    title: str = "Audio Transcription Report",
    transcript_label: str = "Full Transcript (word for word)",
) -> Path:
    """Same order as the results page: transcript, summary, key points, AI insights, highlights."""
    doc = Document()
    doc.add_heading(title, level=0)

    doc.add_heading(transcript_label, level=1)
    doc.add_paragraph(transcript)

    if summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(summary)
    if key_points:
        doc.add_heading("Key Points", level=1)
        _add_bullets(doc, key_points)
    if ai_insights:
        doc.add_heading("AI Insights", level=1)
        doc.add_paragraph(ai_insights)
    if dates_mentioned or links_mentioned or notable_terms:
        doc.add_heading("Highlights", level=1)
        if dates_mentioned:
            doc.add_heading("Dates Mentioned", level=2)
            _add_bullets(doc, dates_mentioned)
        if links_mentioned:
            doc.add_heading("Links Mentioned", level=2)
            _add_bullets(doc, links_mentioned)
        if notable_terms:
            doc.add_heading("Notable Terms (books, courses, names, tools)", level=2)
            _add_bullets(doc, notable_terms)
    return _save(doc)


def build_transcribe_docx(transcript: str) -> Path:
    doc = Document()
    doc.add_heading("Audio Transcript", level=0)
    doc.add_paragraph(transcript)
    return _save(doc)
